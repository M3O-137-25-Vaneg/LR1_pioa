import docx
from fuzzywuzzy import fuzz
import Levenshtein
from time import *

str1 = docx.Document('Doc1.docx')
str2 = docx.Document('Doc2.docx')

str1 = "\n".join([para.text for para in str1.paragraphs])
str2 = "\n".join([para.text for para in str2.paragraphs])

time1 = time()
distance = Levenshtein.distance(str1, str2)
time2 = time()
similarity = Levenshtein.ratio(str1, str2)*100
print(f"Расстояние Левенштейна: {distance}")
print(f"За {time2 - time1}")
print(f"В процентах: {similarity}")

time3 = time()
ratio = fuzz.ratio(str1, str2)
time4 = time()
print(f"В процентах через fuzzy: {ratio}")
print(f"За  {time4 - time3}")

def levenstein(str_1, str_2):
    n, m = len(str_1), len(str_2)
    if n > m:
        str_1, str_2 = str_2, str_1
        n, m = m, n

    current_row = range(n + 1)
    for i in range(1, m + 1):
        previous_row, current_row = current_row, [i] + [0] * n
        for j in range(1, n + 1):
            add, delete, change = previous_row[j] + 1, current_row[j - 1] + 1, previous_row[j - 1]
            if str_1[j - 1] != str_2[i - 1]:
                change += 1
            current_row[j] = min(add, delete, change)

    return current_row[n]

time5 = time()
res = levenstein(str1, str2)
time6 = time()
print(f"Через динамическом программирование levenshtein: {res}")
print(f"За {time6 - time5}")


